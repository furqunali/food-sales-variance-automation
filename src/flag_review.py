#!/usr/bin/env python3
"""
flag_review.py  —  3-month-average deviation flags ("human recheck" gate).

RULE
  For the latest filled month, every tracked value is compared to the AVERAGE of the
  previous up-to-3 months. If it is more than +/-1.5% away from that average, it is
  FLAGGED (a comment/note) so a human can recheck it. This catches typos, missed values,
  and anything unusual in the new source files.

APPROVAL
  A flag stays "pending" (shows a red flag on the dashboard value) until it is approved.
  Once you confirm the numbers are correct, approve them and the flags disappear.

USAGE
  python flag_review.py                      # compute flags for the latest month -> review_flags.json
  python flag_review.py --list               # show all flags and their status
  python flag_review.py --approve all        # approve every pending flag (all correct)
  python flag_review.py --approve 1002 inir  # approve one store+metric
  python flag_review.py --reset              # clear the approvals file entirely

Threshold is THRESHOLD below (0.015 = 1.5%). Flags are stored in review_flags.json next to
the workbook's project so build_dashboard.py can show them on the dashboard.
"""
import os, sys, json, statistics
import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
# The real monthly workbook is private (not in this public repo). See README.
XLSX = os.path.join(ROOT, "data", "food-sales-variance-report.xlsx")
FLAGS = os.path.join(ROOT, "src", "review_flags.json")

THRESHOLD = 0.015          # 1.5%
LOOKBACK = 3               # average of the last 3 months
STORES = [("Cedar Crossing","1001"),("Maple Junction","1002"),("Harbor Point","1003"),("Prairie Gate","1004")]
MONTHS = ["Jan","Feb","Mar","Apr","May","Jun","Jul","Aug","Sep","Oct","Nov","Dec"]

# metric key -> (row base for Jan, column, label, kind)
METRICS = [
    ("actual",   20, 10, "Total sales",      "money"),
    ("cstore",   20,  4, "C-Store sales",     "money"),
    ("kitchen",  20,  7, "Kitchen sales",     "money"),
    ("fountain", 20, 17, "Fountain sales",    "money"),
    ("gp_before", 6,  3, "GP% before disc",   "pct"),
    ("gp_after",  6,  6, "GP% after disc",    "pct"),
    ("inir",      6, 16, "INIR margin",       "pct"),
    ("customers",35,  3, "Customers",         "count"),
    ("sos",       6, 19, "SOS (min)",         "num"),
    ("wastage",   6,  7, "Wastage $",         "money"),
    ("sampling",  6,  9, "Sampling $",        "money"),
    ("logo",      6, 13, "Logo cups $",       "money"),
]

def _n(v):
    return v if isinstance(v, (int, float)) else 0.0

def last_month_idx(ws):
    last = -1
    for i in range(12):
        if _n(ws.cell(20+i, 10).value): last = i
    return last

def fmt(kind, v):
    if kind == "pct":   return f"{v*100:.1f}%"
    if kind == "money": return f"${v:,.0f}"
    if kind == "count": return f"{v:,.0f}"
    return f"{v:.2f}"

def compute():
    wb = openpyxl.load_workbook(XLSX, data_only=True)
    idx = last_month_idx(wb[STORES[0][1]])
    if idx < 0:
        return {}, None
    month = MONTHS[idx]
    flags = {}
    for name, sid in STORES:
        ws = wb[sid]
        for key, base, col, label, kind in METRICS:
            cur = _n(ws.cell(base+idx, col).value)
            prior = [_n(ws.cell(base+idx-k, col).value) for k in range(1, LOOKBACK+1) if idx-k >= 0]
            prior = [p for p in prior if p != 0]
            fid = f"{sid}|{key}|{month}"
            if not prior:
                continue
            avg = statistics.mean(prior)
            if avg == 0:
                if cur != 0:
                    flags[fid] = dict(sid=sid, store=name, metric=key, label=label, kind=kind,
                                      month=month, current=cur, avg3=0, dev=None, dir="new",
                                      note=f"{label}: {fmt(kind,cur)} with no prior baseline - please confirm.")
                continue
            dev = (cur - avg) / avg
            if abs(dev) > THRESHOLD:
                direction = "up" if dev > 0 else "down"
                flags[fid] = dict(sid=sid, store=name, metric=key, label=label, kind=kind,
                                  month=month, current=round(cur, 4), avg3=round(avg, 4),
                                  dev=round(dev, 4), dir=direction,
                                  note=(f"{label} {fmt(kind,cur)} is {abs(dev)*100:.1f}% "
                                        f"{'above' if dev>0 else 'below'} the 3-month avg {fmt(kind,avg)} "
                                        f"- recheck vs source."))
    return flags, month

def load_store():
    if os.path.exists(FLAGS):
        try:
            return json.load(open(FLAGS, encoding="utf-8"))
        except Exception:
            return {}
    return {}

def save_store(d):
    json.dump(d, open(FLAGS, "w", encoding="utf-8"), indent=1)

def refresh():
    """Recompute flags for the latest month, preserving prior approvals by id."""
    computed, month = compute()
    store = load_store()
    old_status = {k: v.get("status", "pending") for k, v in store.items()}
    out = {}
    for fid, f in computed.items():
        f["status"] = old_status.get(fid, "pending")   # keep an earlier approval
        out[fid] = f
    save_store(out)
    return out, month

def show(d):
    pend = [f for f in d.values() if f.get("status") != "approved"]
    appr = [f for f in d.values() if f.get("status") == "approved"]
    if not d:
        print("  No values tracked yet."); return
    print(f"  Flags: {len(pend)} PENDING, {len(appr)} approved  (threshold {THRESHOLD*100:.1f}% vs {LOOKBACK}-month avg)\n")
    for f in sorted(pend, key=lambda x: (x['sid'], x['metric'])):
        print(f"   [PENDING] {f['store']} {f['sid']}  {f['note']}")
    if not pend:
        print("   [OK] No pending flags - all values within 1.5% of trend (or approved).")

def main():
    args = sys.argv[1:]
    if args and args[0] == "--reset":
        if os.path.exists(FLAGS): os.remove(FLAGS)
        print("  review_flags.json cleared."); return
    if args and args[0] == "--list":
        show(load_store()); return
    if args and args[0] == "--approve":
        d = load_store(); tgt = args[1:] or ["all"]
        n = 0
        for fid, f in d.items():
            if f.get("status") == "approved":
                continue
            if tgt == ["all"] or (len(tgt) >= 1 and tgt[0] in fid) and (len(tgt) < 2 or tgt[1] in fid):
                f["status"] = "approved"; n += 1
        save_store(d)
        print(f"  Approved {n} flag(s). Re-run the build to clear them from the dashboard."); return
    # default: recompute for the latest month
    d, month = refresh()
    print(f"=== 3-month-average review flags - month {month} (threshold {THRESHOLD*100:.1f}%) ===")
    show(d)
    print("\n  -> Recheck each PENDING value against its source file.")
    print("  -> If all correct:  python flag_review.py --approve all   (then re-run the build)")

if __name__ == "__main__":
    main()
